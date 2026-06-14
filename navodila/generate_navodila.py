#!/usr/bin/env python3
"""
Generator navodil za Pregled BK.

Prebere NAVODILA.md (omejen Markdown) in posnetke v screens/ ter sestavi
Word datoteko Navodila_PregledBK.docx. Različico aplikacije prebere iz
najvišjega 'build:' v pregled_bk.html, da se navodila ujemajo z aplikacijo.

Zagon:  python generate_navodila.py
Odvisnost:  python-docx  (pip install python-docx)
"""
import re
import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
MD = HERE / "NAVODILA.md"
SCREENS = HERE / "screens"
HTML = HERE.parent / "app" / "src" / "main" / "assets" / "pregled_bk.html"
OUT = HERE / "Navodila_PregledBK.docx"

IMG_WIDTH = Inches(2.7)  # ozki posnetki telefona/čitalnika


def latest_build():
    try:
        txt = HTML.read_text(encoding="utf-8", errors="ignore")
        builds = [int(m) for m in re.findall(r"build:\s*(\d+)", txt)]
        if builds:
            return max(builds)
    except Exception:
        pass
    return None


def add_runs(paragraph, text):
    """Doda besedilo z osnovnim **krepkim** oblikovanjem."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part == "":
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:  # liho = znotraj **...**
            run.bold = True


def add_image(doc, alt, rel):
    img = (HERE / rel).resolve()
    if not img.exists():
        doc.add_paragraph(f"[manjka posnetek: {rel}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=IMG_WIDTH)
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def build():
    doc = Document()
    # privzeta pisava
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = MD.read_text(encoding="utf-8").splitlines()
    img_re = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
    first_h1 = True

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        m = img_re.match(line)
        if m:
            add_image(doc, m.group(1), m.group(2))
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            if first_h1:
                title = doc.add_heading(line[2:], level=0)
                first_h1 = False
                # podnaslov z različico in datumom
                b = latest_build()
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = sub.add_run(
                    ("Različica aplikacije: build %s   ·   " % b if b else "")
                    + "Navodila posodobljena: "
                    + datetime.date.today().strftime("%d.%m.%Y")
                )
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                doc.add_heading(line[2:], level=1)
        elif re.match(r"^\d+\.\s+", line):
            txt = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_runs(p, txt)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

    doc.save(str(OUT))
    print("Ustvarjeno:", OUT)


if __name__ == "__main__":
    build()
